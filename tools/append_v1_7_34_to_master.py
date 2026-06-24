"""Append §12.34 al master doc con notas de v1.7.34+77."""
import hashlib, shutil, sys
from datetime import date
from pathlib import Path
from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SRC = Path(r"G:/Mi unidad/pixoraIA_admin/admin/administracion/Pixora_IA_Documento_Maestro.docx")
SNAP = Path(r"D:/Orbix/Pixora-IA/docs/master_doc_snapshots")
SNAP.mkdir(parents=True, exist_ok=True)
sha8 = hashlib.sha256(SRC.read_bytes()).hexdigest()[:8]
snap_path = SNAP / f"{date.today().isoformat()}_{sha8}.docx"
if not snap_path.exists():
    shutil.copy2(SRC, snap_path)
    print(f"[snapshot] {snap_path.name}")
else:
    print(f"[snapshot] already exists: {snap_path.name}")

doc = Document(str(SRC))

def kv(label, value):
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(value)

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


doc.add_heading("12.34 v1.7.34+77 — 24 de junio 2026", level=2)
kv("Versión", "1.7.34 (versionCode 77)")
kv("Track", "prueba_cerrada")
kv("Fecha", "2026-06-24")
kv("Bump", "1.7.33+76 → 1.7.34+77")

doc.add_paragraph(
    "Release de UX feedback de Eduardo después de usar la app durante "
    "el día. Tres cambios visibles + uno arquitectural que arregla "
    "panorámicos mal detectados."
)

doc.add_heading("Cambios visibles para el usuario", level=3)
bullet(
    "Newest first: las listas (Panorámicos, Anime, Gaming, Calendar, "
    "Mitología, Arte, search, todas las categorías) ahora muestran los "
    "wallpapers MÁS NUEVOS primero. Antes salían los primeros que "
    "subimos (sort_order ASC) y el feed se sentía viejo aunque "
    "hubieran 60+ uploads frescos."
)
bullet(
    "Pill ⟷ PANORAMIC en el holocard del viewer al lado del ★★★ RARE "
    "cuando el wallpaper es panorámico. Magenta brillante (#FF2BD6). "
    "User identifica de un vistazo qué tipo va a aplicar."
)
bullet(
    "Scanline del viewer ya NO es permanente. Antes la línea recorría "
    "la imagen para siempre (parecía 'imagen cargando' eterno). Ahora "
    "es una sola pasada de 2.5s (scan-and-reveal) y se queda quieta. "
    "Re-dispara cada vez que swipe a otra imagen — efecto de "
    "bienvenida visual sin distracción permanente."
)

doc.add_heading("Cambios técnicos", level=3)
bullet(
    "Wallpaper.type field: nuevo campo opcional que mapea el enum "
    "wallpaper_type de Postgres (static / panoramic / live / canvas_scene). "
    "Wallpaper.fromSupabase y fromJson lo leen. isPanoramic ahora revisa "
    "3 señales en orden: type == 'panoramic' (explicit) → ratio >= 3.0 "
    "(ultra-wide auto-detect) → category == 'PANORAMIC' (legacy)."
)
bullet(
    "sort_order DESC: cambiado en CatalogService (línea 132 + 193) + "
    "8 lugares en wallpaper_providers.dart. CatalogService order desde "
    "Postgres + sort() en el parser JSON + sort() en cada provider de "
    "sección. Pattern: `b.sortOrder.compareTo(a.sortOrder)`."
)
bullet(
    "Backend de la misma fecha (no requiere v1.7.34 actualizada): "
    "3 panorámicas que vivían mal clasificadas se movieron a "
    "category=PANORAMIC en Postgres: pano_adventure_time_noche, "
    "pano_coyote_correcaminos, pano_pixel_palace_arcade. v1.7.30+ "
    "actual también se beneficia."
)

doc.add_heading("Lecciones aprendidas", level=3)
bullet(
    "Defensa en profundidad para detección de tipo: cuando un campo "
    "puede vivir en múltiples lugares (DB type + DB category + "
    "aspectRatio de las dimensiones), TODO check de detección debe "
    "revisar las 3 señales, no solo una. Antes solo revisaba ratio + "
    "category, así que cualquier panorámico con ratio < 3 y category "
    "no-PANORAMIC se volvía invisible al cliente."
)
bullet(
    "sort_order ASC vs DESC: la convención inicial de poner los más "
    "viejos primero (ASC) tiene sentido para catálogos archivo (history) "
    "pero NO para feeds (queremos surface freshness). Para cualquier "
    "feed de consumidor: DESC siempre. Bug encubierto hasta que "
    "Eduardo notó que sus uploads recientes no aparecían al inicio."
)
bullet(
    "AnimationController.repeat() es UX-hostil para overlays no-essenciales. "
    "Un efecto 'sweep' debería ser one-shot (forward + stop) o gated por "
    "loading state. Permanente da sensación de 'algo está mal o cargando'. "
    "Pattern: `.forward()` + `.forward(from: 0)` en cambio de página."
)

doc.save(str(SRC))
print("[OK] §12.34 appended")
