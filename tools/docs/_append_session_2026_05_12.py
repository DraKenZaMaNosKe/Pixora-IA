"""
End-of-session master doc append for 2026-05-12.

Appends:
  - §12 v1.7.13+56 release entry
  - §19+ "Reporte de progreso · 12 Mayo 2026" with FCM + Dashboard CRUD + v1.7.13

Snapshots first to docs/master_doc_snapshots/ (gitignored).
Redaction: no secret values pasted — KEYS_LOCAL.md references only.
"""
from __future__ import annotations
import sys, hashlib, shutil
from datetime import date
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document  # noqa: E402

SRC = Path(r"G:/Mi unidad/pixoraIA_admin/admin/administracion/Pixora_IA_Documento_Maestro.docx")
SNAP_DIR = Path(r"D:/Orbix/Pixora-IA/docs/master_doc_snapshots")

# 1) snapshot
SNAP_DIR.mkdir(parents=True, exist_ok=True)
sha8 = hashlib.sha256(SRC.read_bytes()).hexdigest()[:8]
snap = SNAP_DIR / f"{date.today().isoformat()}_{sha8}.docx"
if not snap.exists():
    shutil.copy2(SRC, snap)
    print(f"Snapshot saved: {snap.name}")
else:
    print(f"Snapshot for {sha8} already exists")

doc = Document(str(SRC))


def kv(label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(value)


def bullet(text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


# ===================================================================
# §12 RELEASE LOG — v1.7.13+56
# ===================================================================
doc.add_heading("12.X v1.7.13+56 — 2026-05-12 — Push notifications + Apple Music preview", level=2)
kv("Fecha", "2026-05-12 01:00 CST")
kv("Track", "Closed Testing (segmento 'prueba_cerrada')")
kv("AAB", "build/app/outputs/bundle/release/app-release.aab (134.4 MB)")
kv("AAB SHA-1 firma", "FF:0F:46:D4:E2:86:25:D1:14:C6:81:03:11:E4:6B:E4:47:2A:CD:79")
kv("Commit version bump", "f169839")
kv("Git tag", "v1.7.13 (local — push pendiente)")
doc.add_paragraph().add_run("Cambios shippeados en esta versión:").bold = True
bullet("FCM Push Notifications — usuario recibe push cuando publicamos wallpapers nuevos. Topic broadcast 'new_content'.")
bullet("Apple Music 'Now Playing' redesign del preview page — preview tipo álbum + texto legible (cream bg). Concept #01 elegido de 5 diseños propuestos.")
bullet("sortOrder fix — wallpapers más nuevos aparecen primero (sort por createdAt DESC, fallback sortOrder DESC).")
bullet("effectiveBadge auto-expira NEW después de 7 días (en lib/features/hot_wallpapers/data/models/live_wallpaper.dart).")
bullet("3 wallpapers nuevos LIVE: Café Tarde Lluviosa (ANIME), Claire Redfield (GAMING), NYC Outbreak (GAMING).")
doc.add_paragraph().add_run("Notas técnicas:").bold = True
bullet("AdMob: TEST AD ID sigue activo (intencional hasta llegar a Production track).")
bullet("Firebase service account JSON respaldado en orbixprivate/secrets/pixora-firebase/.")
bullet("Build inicial falló (Gradle exit 1 silencioso tras 34 min) — segunda corrida OK (13 min). Causa raíz no identificada pero reproducible: relanzar sin tocar nada lo arregla.")

# ===================================================================
# §19+ REPORTE DE PROGRESO — 12 Mayo 2026
# ===================================================================
doc.add_heading("Reporte de Progreso · 12 Mayo 2026 · FCM + Dashboard CRUD + v1.7.13", level=2)
kv("Fecha", "2026-05-12 (sesión nocturna ~3 horas)")
kv("Tipo de sesión", "Cierre de features + release prep + dashboard tooling")
kv("Estado al iniciar", "v1.7.12+55 publicada el día anterior, FCM ya integrado, faltaba subir 3 wallpapers nuevos y redesign del preview page")
kv("Estado al cerrar", "v1.7.13+56 lista para subir a Play Console, dashboard CRUD funcional para 5 catálogos")

doc.add_paragraph().add_run("Objetivos cumplidos:").bold = True

bullet("Implementado concept #01 (Apple Music Now Playing) en live_wallpaper_preview_page.dart. Cream bg #FAF7F2, preview tipo álbum-art, Fraunces italic title, controls row (♡⤴ⓘ), CTA azul iOS rounded.")

bullet("Subidos 3 wallpapers LIVE a producción: Café Tarde Lluviosa, Claire Redfield, NYC Outbreak. Visible para todos sin actualizar APK porque viven en catálogo Storage.")

bullet("Bump version 1.7.12 → 1.7.13+56. AAB construido, firmado, verificado contra upload key. Release notes ES/EN <500 chars cada una en docs/release_notes/v1.7.13.md.")

bullet("Dashboard CRUD wallpapers (Pixora Admin local). Modal de cualquier wallpaper en LIVE o STATIC catalog tiene botón Editar — permite cambiar name, description, category, sortOrder, badge, glowColor, tags. También eliminar del catálogo. Cambios escriben el JSON al instante a Supabase Storage. Útil porque Eduardo quiere curar el catálogo a mano sin tocar SQL.")

bullet("Arquitectura del CRUD: server local (wp_admin_server.py) proxea GET/PUT /api/catalog/<kind> usando SERVICE_KEY que ya tenía cargada al startup. El browser nunca ve la key. Soporta 5 catálogos: live, static, stories, day_cycle, ringtones (los últimos 3 wired pero sin items aún en el grid del dashboard).")

bullet("Dashboard responsive fixes: stack vertical a <900px (antes 720), overflow-x:hidden en panels del modal, stat-grid con auto-fit minmax(88px,1fr) en vez de repeat(4,1fr) fijo, min-width:0 en grid tracks. Mata el scroll horizontal feo en laptops chicas.")

bullet("Standalone dashboard_crud.html — versión sola del CRUD que no requiere wp_admin_server. Lee/escribe Storage directo con SERVICE_KEY desde dashboard_crud_key.local.js (gitignored). Útil cuando el admin_server está caído o se ejecuta desde laptop sin Python.")

doc.add_paragraph().add_run("Pendientes para próxima sesión (oficina):").bold = True
bullet("Subir AAB v1.7.13+56 a Play Console (Closed Testing) con las release notes ya redactadas.")
bullet("Subir wallpaper Monkey D. Luffy — usuario tiene MonkeyD.png (1080x2160 static) + posiblemente video MP4 — pendiente confirmar ruta del video. Categoría ANIME, título 'Monkey D. Luffy', free con ads.")
bullet("Push tag v1.7.13 y commits a origin (no se hizo en esta sesión — pendiente revisión).")
bullet("Probar el CRUD del dashboard en producción editando un wallpaper real (ej. corregir typo en alguna descripción).")
bullet("Wire de stories/day_cycle/ringtones al grid del dashboard (server ya soporta los catálogos, pero el grid solo lista live+static).")

doc.add_paragraph().add_run("Decisiones de diseño/producto:").bold = True
bullet("'No admin UI dentro de Pixora app' confirmado otra vez — el CRUD de catálogo vive en Pixora Admin (localhost), nunca en el APK distribuible. Usuario lo curará a mano antes de cada release.")
bullet("Preview page redesign — Eduardo eligió concept #01 (Apple Music Now Playing) de 5 propuestas. Las otras 4 (Translucent Glass, Dark Gradient, Split View, Apple TV+ Hero) quedan archivadas en docs/design/wallpaper_preview_page_ios_concepts.html.")

# ===================================================================
# Save
# ===================================================================
doc.save(str(SRC))
print("Master doc actualizado.")
print(f"  Snapshot: {snap.name}")
print("  §12 entry: v1.7.13+56")
print("  §19 entry: Reporte de Progreso 12 Mayo 2026")
