"""
End-of-session master doc append for 2026-05-14.

Append:
  - §12 v1.7.15+58 release entry
  - §19+ "Reporte de progreso 14 Mayo 2026"

Snapshot first to docs/master_doc_snapshots/.
Redact secrets: reference KEYS_LOCAL.md only.
"""
from __future__ import annotations
import sys, hashlib, shutil
from datetime import date
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document  # noqa: E402

SRC = Path(r"G:/Mi unidad/pixoraIA_admin/admin/administracion/Pixora_IA_Documento_Maestro.docx")
SNAP_DIR = Path(r"D:/Orbix/Pixora-IA/docs/master_doc_snapshots")

SNAP_DIR.mkdir(parents=True, exist_ok=True)
sha8 = hashlib.sha256(SRC.read_bytes()).hexdigest()[:8]
snap = SNAP_DIR / f"{date.today().isoformat()}_{sha8}.docx"
if not snap.exists():
    shutil.copy2(SRC, snap)
    print(f"Snapshot saved: {snap.name}")
else:
    print(f"Snapshot for {sha8} already exists")

doc = Document(str(SRC))


def kv(label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(value)


def bullet(text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


# ===================================================================
# §12 RELEASE LOG — v1.7.15+58
# ===================================================================
doc.add_heading("12.X v1.7.15+58 — 2026-05-14 — Pixora Daily curado + Arte/Mitologia sections", level=2)
kv("Fecha", "2026-05-14 (sesion oficina-noche)")
kv("Track", "Closed Testing (segmento 'prueba_cerrada')")
kv("AAB", "build/app/outputs/bundle/release/app-release.aab")
kv("AAB SHA-1 firma", "FF:0F:46:D4:E2:86:25:D1:14:C6:81:03:11:E4:6B:E4:47:2A:CD:79")
kv("Commits desde v1.7.14", "e8c1d11 + cd68352 + bump 1.7.15")
kv("Git tag", "v1.7.15 (local pendiente push)")

doc.add_paragraph().add_run("Cambios shippeados:").bold = True
bullet("Pixora Daily curado — admin marca wallpapers como daily_eligible (boolean) en dashboard, el AutoRotateService los filtra cuando usuario elige 'Pixora Daily (curado)' en el picker. Mantiene categoria natural del wallpaper. Resuelve la queja de mezcla rara (calendario junto a anime).")
bullet("Secciones por tags — wallpapers_page.dart agrega 2 carousels nuevos: Arte y Mitologia. Filtran por tag 'arte'/'mitologia' OR categoria match. Multi-categoria real sin migrar enum a array.")
bullet("Categoria PANORAMIC + tag 'arte' = wallpaper aparece en seccion PANORAMIC Y en seccion Arte. Ejemplo Surreal Dali Dreamscape funciona en ambos lados.")
bullet("Fix overflow 74px en Pixora Daily picker — Flexible + SingleChildScrollView + ConstrainedBox(maxHeight 75%). isScrollControlled en showModalBottomSheet.")

doc.add_paragraph().add_run("Database:").bold = True
bullet("Columna wallpapers.daily_eligible boolean + partial index para filtros rapidos")
bullet("Vista wallpapers_v recreada exponiendo daily_eligible")
bullet("Enum wallpaper_category: +DAILY +CALENDARIOS +MITOLOGIA. Total 31 valores. DAILY/CALENDARIOS quedan dormant en UI (DAILY se reemplazo por flag, CALENDARIOS por error del admin que admitio fue duplicado de CALENDAR existente)")

doc.add_paragraph().add_run("Dashboard Pixora Admin:").bold = True
bullet("Toggle 'Eligible para Pixora Daily' en modal de edit, separado del toggle Publicado")
bullet("Badge estrella azul-celeste en cards daily-eligible")
bullet("Lightbox full-res: click en preview del modal -> imagen al 92vw x 92vh, ESC para cerrar")
bullet("Hint de dimensiones (naturalWidth/Height) + orientacion + ratio")
bullet("Warning rojo si admin elige categoria incompatible (PANORAMIC en 9:16, AURA en no-audio). No bloquea guardado, solo avisa.")
bullet("MITOLOGIA agregada al dropdown grupo 'Tuyas'")

# ===================================================================
# §19+ REPORTE DE PROGRESO — 14 Mayo 2026
# ===================================================================
doc.add_heading("Reporte de Progreso 14 Mayo 2026 — Multi-categoria + curaduria Pixora Daily + UX fixes", level=2)
kv("Fecha", "2026-05-14")
kv("Estado al iniciar", "v1.7.14 publicada y firmada en Play Store Closed Testing. Bug 22P02 en dashboard al guardar categoria (CHILL no existia en enum, ya rescatamos). Pixora Daily mezclaba wallpapers de categorias dispares.")
kv("Estado al cerrar", "v1.7.15 lista para subir. Multi-categoria via tags funcionando. Curaduria Daily independiente de categoria. UI overflow corregido. Validacion preventiva en dashboard.")

doc.add_paragraph().add_run("Decisiones arquitecturales:").bold = True
bullet("Multi-categoria via TAGS, no array de categorias. category queda como funcion estructural single-value (PANORAMIC, AURA, etc.) y tags como discovery semantico multi-value (arte, mitologia, daily). Documentado en memoria persistente tech_category_vs_tags_pixora.md.")
bullet("daily_eligible como BOOLEAN flag, no como categoria. Asi el wallpaper conserva su categoria natural (ANIME, GAMING) y tambien puede rotar en Daily curado. Patron similar a published flag.")
bullet("Convencion de tags-sección: lowercase sin acentos (arte, mitologia, daily). Tags descriptivos con guiones (dragon-ball, pixel-art).")
bullet("Nivel 2 de validacion preventiva en dashboard: hint de dimensiones + warning si elige categoria incompatible. No bloquea (Nivel 3 era server-side reject — descartado por false positives). Decision del usuario al final.")

doc.add_paragraph().add_run("Pendientes proxima sesion:").bold = True
bullet("Subir AAB v1.7.15+58 a Play Console Closed Testing")
bullet("Empezar a marcar wallpapers de arte como tag 'arte' (Dali, museos, gallery wallpapers existentes en static catalog)")
bullet("Marcar wallpapers de mitologia con tag 'mitologia' (Mictlantecuhtli, Iah-related, etc.)")
bullet("Curar inicialmente 5-10 wallpapers como daily_eligible para que Pixora Daily curado tenga pool de rotacion")
bullet("Conseguir 12 testers para Closed Testing — meta usuario: fin de mayo")
bullet("Eventual: subir Monkey D. Luffy version static (LIVE ya subido)")

doc.save(str(SRC))
print("Master doc actualizado.")
print(f"  Snapshot: {snap.name}")
print("  §12 entry: v1.7.15+58")
print("  §19 entry: Reporte de Progreso 14 Mayo 2026")
