"""
End-of-session master doc append for 2026-05-13.

Append:
  - §12 v1.7.14+57 release entry
  - §19+ "Reporte de progreso 13 Mayo 2026" — performance + dashboard CRUD

Snapshot first to docs/master_doc_snapshots/.
Redact secrets: reference KEYS_LOCAL.md only.
"""
from __future__ import annotations
import sys, hashlib, shutil
from datetime import date
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document  # noqa: E402

SRC = Path(r"G:/Mi unidad/pixoraIA_admin/admin/administracion/Pixora_IA_Documento_Maestro.docx")
SNAP_DIR = Path(r"D:/Orbix/Pixora-IA/docs/master_doc_snapshots")

# 1) snapshot
SNAP_DIR.mkdir(parents=True, exist_ok=True)
sha8 = hashlib.sha256(SRC.read_bytes()).hexdigest()[:8]
snap = SNAP_DIR / f"{date.today().isoformat()}_{sha8}.docx"
if not snap.exists():
    shutil.copy2(SRC, snap)
    print(f"Snapshot saved: {snap.name}")
else:
    print(f"Snapshot for {sha8} already exists")

doc = Document(str(SRC))


def kv(label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(value)


def bullet(text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


# ===================================================================
# §12 RELEASE LOG — v1.7.14+57
# ===================================================================
doc.add_heading("12.X v1.7.14+57 — 2026-05-13 — Cold-start performance + Aurora UI", level=2)
kv("Fecha", "2026-05-13 (sesion vespertina-nocturna)")
kv("Track", "Closed Testing (segmento 'prueba_cerrada')")
kv("AAB", "build/app/outputs/bundle/release/app-release.aab")
kv("AAB SHA-1 firma", "FF:0F:46:D4:E2:86:25:D1:14:C6:81:03:11:E4:6B:E4:47:2A:CD:79")
kv("Commit version bump", "0a9aac6")
kv("Git tag", "v1.7.14 (local, push pendiente)")

doc.add_paragraph().add_run("Cambios shippeados:").bold = True
bullet("Cold-start instantaneo — CatalogService.preloadFromDiskCache() + LiveWallpaperCatalogService.preloadFromDiskCache() leen el cache de disco antes de runApp(). Resultado: tab Wallpapers/LIVE arrancan con cards al instante en vez de skeleton 1-3s.")
bullet("LIVE tab refactorizada a CustomScrollView + SliverGrid por categoria. Antes era SingleChildScrollView+Column+GridView(shrinkWrap,NeverScrollable) que construia TODO al inicio — bug latente que pegaria fuerte al pasar de ~50 wallpapers por categoria.")
bullet("AuroraWavesLoading widget reemplaza Shimmer.fromColors generico. CustomPainter con 4 blobs radiales que ondulan en curva Lissajous + shimmer vertical. Theme-aware: B&G usa morados/dorados profundos, iOS pasteles rosa/lavanda/menta/durazno. Concept #01 elegido de 5 mockups HTML.")
bullet("Catalog loading skeleton tambien con Aurora Waves dentro de card shapes (Faithful Mirror concept #01 de 5). Adios shimmer Material generico.")
bullet("LiveWallpaperCatalogService gana persistencia a disco — live_wallpaper_cache.json guarda response raw. Robusto ante cambios de schema del catalogo.")
bullet("Cache TTL: 6h -> 30min en ambos servicios. Cambios del dashboard llegan al usuario 12x mas rapido.")
bullet("Pull-to-refresh ahora llama clearCache() del singleton antes de invalidar Riverpod — fix del bug donde el TTL impedia el refetch real.")

doc.add_paragraph().add_run("Dashboard Pixora Admin:").bold = True
bullet("POST /api/wallpaper-edit nuevo endpoint que actualiza Postgres `wallpapers` table + Storage JSON en una sola operacion atomica. Fix del bug donde edits del dashboard solo escribian al JSON y la app no veia cambios (la app lee de Postgres, no de Storage para STATIC).")
bullet("Modal del wallpaper ahora lee del catalogo JSON (source of truth real) en vez de admin_wallpaper_breakdown view (que tenia nombres cacheados de eventos antiguos).")
bullet("4 categorias nuevas en el dropdown: PAISAJISMO, TV, ARTE, VARIOS.")
bullet("Scripts one-shot tools/wallpapers/_reconcile_static_catalog.py y _reverse_sync_pg_to_storage.py para sincronizar Storage<->Postgres tras desfases.")

# ===================================================================
# §19+ REPORTE DE PROGRESO — 13 Mayo 2026
# ===================================================================
doc.add_heading("Reporte de Progreso 13 Mayo 2026 — Performance audit + UI consistency + dashboard CRUD complete", level=2)
kv("Fecha", "2026-05-13 (sesion ~7 horas, oficina)")
kv("Estado al iniciar", "v1.7.13+56 publicada y firmada localmente. AAB listo pero no subido. CRUD dashboard funcionando pero solo escribia al JSON Storage, no a Postgres. Pull-to-refresh con bug (TTL impedia refetch). Loading state era Shimmer generico Material.")
kv("Estado al cerrar", "v1.7.14+57 lista para subir. Dashboard CRUD escribe a ambos lados (atomic). Cold-start instantaneo. UI loading con Aurora Waves consistente. Memoria persistente de workflow de skills para futuras sesiones.")

doc.add_paragraph().add_run("Objetivos cumplidos:").bold = True

bullet("Audit completo de performance del catalog flow. Identificadas 3 dimensiones de bottleneck: (1) tamaño del JSON catalog, (2) grids no-lazy en LIVE tab, (3) cold-start sin preload de disco. Las 3 atacadas en esta sesion. Documentado threshold honesto: hoy 342 wallpapers sin problemas, fixes preventivos te dejan margen hasta ~10,000 wallpapers totales sin necesidad de pagination.")

bullet("Fix critical: Dashboard CRUD escribia solo al JSON Storage pero la app (para STATIC wallpapers) lee de Postgres `wallpapers` tabla. Resultado: edits en dashboard no llegaban al app. Solucion: endpoint POST /api/wallpaper-edit en el server local que hace UPDATE Postgres + PUT Storage en una sola operacion.")

bullet("One-time reconciliation: 66 wallpapers editados antes del fix fueron sincronizados Storage -> Postgres via script _reconcile_static_catalog.py. 2 wallpapers fallaron (constraint NOT NULL en sort_order, no critico). Posteriormente se hizo reverse sync Postgres -> Storage para limpiar 2 diferencias residuales.")

bullet("Nuevo widget AuroraWavesLoading. Aprobado por el usuario despues de mostrar 5 conceptos HTML (Aurora Waves, Constellations, Gold Shimmer, Pulsing Sigil, Cinematic Iris). Eduardo escogio NO.01 Aurora Waves. Implementado en Flutter con CustomPainter performante: 4 radial gradient blobs que driftean via curva Lissajous + shimmer vertical. ~1ms/frame por card en Samsung gama media.")

bullet("Catalog loading skeleton tambien con Aurora Waves (Faithful Mirror concept #01 de otros 5 mockups HTML). _ShimmerLoading y CarouselRowShimmer rebuild para usar Aurora dentro de card shapes en lugar de Shimmer.fromColors.")

bullet("hot_wallpapers_page.dart refactor de SingleChildScrollView a CustomScrollView+SliverGrid. Antes el shrinkWrap forzaba construir TODOS los items al inicio. Ahora truly lazy: solo se construye lo visible. Escala a miles de wallpapers sin trabarse.")

bullet("Pull-to-refresh implementado correctamente en ambas tabs (estatica y LIVE). Llama clearCache() del singleton antes de invalidar Riverpod para forzar refetch real. Previamente el TTL bloqueaba la operacion.")

bullet("Cold-start preload: CatalogService.preloadFromDiskCache() + LiveWallpaperCatalogService.preloadFromDiskCache() ejecutan en paralelo (Future.wait) antes de runApp(). Resultado en device test: la primera vez tarda 1-3s (no hay cache aun), siguientes cold-starts arrancan instantaneos.")

bullet("LiveWallpaperCatalogService gano persistencia a disco. Guarda response.body raw en live_wallpaper_cache.json en cada fetch exitoso. Mas robusto que serializar el modelo Dart porque preserva campos nuevos del JSON aunque no esten en el modelo.")

bullet("Cache TTL bajado de 6h a 30min en ambos catalog services. Los edits del dashboard llegan al usuario en max 30 min sin necesidad de FCM force-refresh.")

bullet("Memoria persistente feedback_skills_workflow_pixora.md documentando que skills usar y cuando — para que futuras sesiones de Claude auto-apliquen el setup correcto sin que el usuario tenga que memorizar reglas.")

doc.add_paragraph().add_run("Decisiones de diseno/producto:").bold = True
bullet("Aurora Waves elegido para loading de imagenes individuales (concept #01 de 5).")
bullet("Faithful Mirror skeleton elegido para loading de catalog completo (concept #01 de 5).")
bullet("Cache TTL 30min mas pull-to-refresh manual aceptado como suficiente (en lugar de FCM force-refresh data messages). Razon: 50% chance de exito en primer intento por la complejidad de background isolates + Doze mode + iOS APNs. FCM puede venir en v1.7.15+ si llega a hacer falta.")
bullet("Dashboard accepted 4 categorias nuevas: PAISAJISMO, TV, ARTE, VARIOS — para mejor organizacion del catalogo creciente.")

doc.add_paragraph().add_run("Pendientes proxima sesion:").bold = True
bullet("Subir AAB v1.7.14+57 a Play Console Closed Testing.")
bullet("Push tag v1.7.14 y commits a origin (pendiente confirmacion del usuario).")
bullet("Subir Monkey D. Luffy version static (PNG -> WebP -> wallpaper-images bucket + entry en dynamic_catalog.json). Hoy solo se subio LIVE version.")
bullet("Monitorear logcat en device tras release a Production para detectar regresiones de performance (especialmente en cold-start sin cache previo).")
bullet("Evaluar deuda tecnica: el `_loadFromCache()` de catalog_service.dart aun se usa como step 3 fallback, pero ahora con preloadFromDiskCache() es redundante. Refactor posible en sesion futura.")

# ===================================================================
# Save
# ===================================================================
doc.save(str(SRC))
print("Master doc actualizado.")
print(f"  Snapshot: {snap.name}")
print("  §12 entry: v1.7.14+57")
print("  §19 entry: Reporte de Progreso 13 Mayo 2026")
